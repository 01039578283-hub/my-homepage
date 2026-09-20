"""Verify original center sources, stage Step 6, then optionally apply locally.

No repository, hosting or Search Console writes. Run without --write first.
Private reports/backups are outside the static site and deployment tree.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from collections import Counter
from pathlib import Path

from lxml import html
from openpyxl import load_workbook

from branch_course_guidance import course_answer, course_guidance, verify_source
from branch_learning_routes import ROOT, canonical, paired_routes
from generate_branch_directory import SOURCE_WORKBOOK, SUPPLEMENTAL_CENTERS_FILE, clean, split_values
from subject_guide_facts import (
    REVIEW_DATE, STYLE, TARGETS_FILE, align_guide, condition_texts, corrected_json,
    exact_fact_substitutions, factual_question, source_schools, targets,
)

REPORT_ROOT = Path(r'C:\Users\1992k\Desktop\CodexData\audit-output\wawa-center-step6-20260920')
ORG_FIELDS = ('address', 'legalName', 'url', 'identifier', 'openingHoursSpecification', 'openingHours', 'makesOffer', 'offers', 'educationalLevel', 'teaches')


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def graph(doc) -> list[dict]:
    return json.loads(doc.xpath('//script[@type="application/ld+json"]/text()')[0])['@graph']


def verify_centers(records: list[dict], centers: dict) -> list[dict]:
    verify_source(SOURCE_WORKBOOK)
    config = json.loads(TARGETS_FILE.read_text(encoding='utf-8'))
    if sha(SOURCE_WORKBOOK) != config['sourceWorkbookSha256']:
        raise ValueError('대상 검토본과 엑셀 원본이 다릅니다.')
    workbook = load_workbook(SOURCE_WORKBOOK, read_only=True, data_only=True)
    sheet = workbook['센터정보']
    source_rows = {i: row for i, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2)}
    supplemental = {c['routeName']: c for c in json.loads(SUPPLEMENTAL_CENTERS_FILE.read_text(encoding='utf-8'))['centers']}
    evidence = []
    for name in sorted({r['center'] for r in records}):
        center = centers[name]
        if center['sourceRow']:
            number = center['sourceRow']
            row = source_rows[number]
            for key, index in (('sourceName', 0), ('registeredName', 6), ('registrationNumber', 7), ('address', 11)):
                if center[key] != clean(row[index]):
                    raise ValueError(f'센터 원본 불일치: {name}/{key}')
            for subject, index in (('국어', 16), ('영어', 17), ('수학', 18), ('과학', 19), ('사회', 20)):
                if center['subjects'].get(subject, []) != split_values(row[index]):
                    raise ValueError(f'학년 원본 불일치: {name}/{subject}')
            brand = '모두오름학습코칭학원' if '(모두)' in clean(row[0]) or '모두오름' in clean(row[6]) else '와와학습코칭센터'
            if center['displayName'] != f'{brand} {name}':
                raise ValueError(f'브랜드 원본 불일치: {name}')
            evidence.append({'center': name, 'source': SOURCE_WORKBOOK.name, 'sheet': sheet.title, 'cells': [f'A{number}', f'G{number}:H{number}', f'L{number}', f'Q{number}:U{number}', f'V{number}', f'AE{number}']})
        else:
            if name not in supplemental or not center.get('sourceProvenance'):
                raise ValueError(f'별도 확인 자료 누락: {name}')
            for key in ('address', 'subjects', 'registeredName', 'registrationNumber', 'brand', 'region', 'district'):
                if center[key] != supplemental[name][key]:
                    raise ValueError(f'별도 확인 자료 불일치: {name}/{key}')
            evidence.append({'center': name, 'source': SUPPLEMENTAL_CENTERS_FILE.name, 'reviewedAt': center['sourceProvenance']['reviewedAt']})
    workbook.close()
    return evidence


def scope_fingerprint(raw: str, record: dict, center: dict) -> dict:
    doc = html.document_fromstring(raw)
    exact_fact_substitutions(doc, record, center)
    schema = corrected_json(graph(doc), record, center)
    view = course_guidance(center, record['subject'], '고')
    if not view['grades']:
        schema = [n for n in schema if n.get('@type') != 'Service']
    school_names = set(record['previousSchools']) | set(source_schools(center))
    for node in schema:
        kind = node.get('@type')
        if kind in ('EducationalOrganization', 'LocalBusiness'):
            for field in ORG_FIELDS:
                node.pop(field, None)
        if kind == 'WebPage':
            node.pop('mainEntity', None)
        if kind in ('WebPage', 'Article'):
            node.pop('dateModified', None)
        if kind == 'Service':
            for field in ('description', 'audience', 'makesOffer', 'offers'):
                node.pop(field, None)
        if kind in ('Article', 'WebPage', 'Service'):
            node['mentions'] = [n for n in node.get('mentions', []) if not isinstance(n, dict) or n.get('name') not in school_names]
        if kind == 'FAQPage':
            node['mainEntity'] = [entry for entry in node['mainEntity'] if entry.get('name') != factual_question(center, record)]
    for element in doc.xpath(f'//link[@href="{STYLE}"] | //*[@id="guide-course-note" or @id="verified-course-faq"] | //script[@type="application/ld+json"]'):
        element.getparent().remove(element)
    card = doc.xpath('//aside[@class="math-info-card"]')[0]
    card.attrib.pop('id', None)
    card.attrib.pop('data-facts-reviewed', None)
    labels = {'센터 기준', '등록 명칭', '제공 주소', '교육지원청 등록번호', '제공 학교 참고', '고등학교 참고', '정보 확인 기준일', '확인할 수업 조건', f'고등 {record["subject"]} 수업 가능 학년', f'고등 {record["subject"]} 안내 학년'}
    for row in list(card.find('dl')):
        if row.find('dt') is not None and row.find('dt').text_content() in labels:
            card.find('dl').remove(row)
    for element in doc.iter():
        if element.text is not None and not element.text.strip():
            element.text = None
        if element.tail is not None and not element.tail.strip():
            element.tail = None
    return {'schema': schema, 'html': html.tostring(doc, encoding='unicode')}


def factual_errors(raw: str, record: dict, center: dict) -> list[str]:
    errors = []
    doc = html.fromstring(raw)
    schema = graph(doc)
    card = doc.xpath('//*[@id="reviewed-center-info"]')
    if len(card) != 1:
        return ['missing reviewed center card']
    rows = {row.get('data-guide-fact'): row.find('dd') for row in card[0].find('dl') if row.get('data-guide-fact')}
    fields = {'name': center['displayName'], 'legal-name': center['registeredName'], 'address': center['address'], 'registration': center['registrationNumber']}
    for key, value in fields.items():
        if key not in rows or rows[key].text_content() != value:
            errors.append(f'visible {key}')
    view = course_guidance(center, record['subject'], '고')
    if rows['grades'].xpath('.//div[@class="math-tag-list"]/span/text()') != view['grades']:
        errors.append('visible grades')
    if not view['grades'] and rows['grades'].text_content() != '개설 학년 문의':
        errors.append('unconfirmed grade wording')
    if [p.text_content() for p in rows['conditions'].findall('p')] != condition_texts(center, record):
        errors.append('visible course conditions')
    if rows['schools'].xpath('.//span/text()') != source_schools(center):
        errors.append('visible school level')
    if center['informationReviewedAt'] not in rows['review-date'].text_content():
        errors.append('source review date')
    answer = course_answer(center, record['subject'], '고')
    faq = doc.xpath('//*[@id="verified-course-faq"]')
    if len(faq) != 1 or faq[0].find('summary').text_content() != factual_question(center, record) or faq[0].find('p').text_content() != answer:
        errors.append('factual FAQ')
    notices = doc.xpath('//*[@id="guide-course-note"]')
    if (not view['grades'] or view['pendingGrades']):
        if len(notices) != 1 or answer not in notices[0].text_content():
            errors.append('missing early qualification')
    elif notices:
        errors.append('unnecessary qualification')
    if len(doc.xpath(f'//link[@href="{STYLE}"]')) != 1:
        errors.append('stylesheet')
    orgs = [n for n in schema if n.get('@type') in ('EducationalOrganization', 'LocalBusiness')]
    if len(orgs) != 2:
        errors.append('organization count')
    expected_parent = canonical(paired_routes(center, record)['parent'])
    for node in orgs:
        if node['name'] != center['displayName'] or node.get('legalName') != center['registeredName'] or node['address']['streetAddress'] != center['address'] or node['identifier']['value'] != center['registrationNumber'] or node['url'] != expected_parent:
            errors.append('organization facts')
        if any(field in node for field in ('openingHoursSpecification', 'openingHours', 'makesOffer', 'offers', 'educationalLevel', 'teaches')):
            errors.append('unsupported organization facts')
    services = [n for n in schema if n.get('@type') == 'Service']
    if len(services) != int(bool(view['grades'])):
        errors.append('unsupported course service')
    for service in services:
        if service['audience']['audienceType'] != view['label'] or service['description'] != answer:
            errors.append('service audience or conditions')
        if 'makesOffer' in service or 'offers' in service:
            errors.append('unverified course offer')
        if service['provider']['@id'] not in {n['@id'] for n in orgs}:
            errors.append('dangling provider')
    page = next(n for n in schema if n.get('@type') == 'WebPage')
    article = next(n for n in schema if n.get('@type') == 'Article')
    if page['mainEntity'] != {'@id': article['@id']}:
        errors.append('learning guide mainEntity')
    if not view['grades'] and '#service' in json.dumps(schema, ensure_ascii=False):
        errors.append('dangling unavailable service reference')
    visible_faqs = [(' '.join(d.find('summary').text_content().split()), ' '.join(d.find('p').text_content().split())) for d in doc.xpath('//div[@class="math-faq-list"]/details')]
    schema_faqs = [(q['name'], q['acceptedAnswer']['text']) for q in next(n for n in schema if n.get('@type') == 'FAQPage')['mainEntity']]
    if visible_faqs != schema_faqs or len({q for q, _ in visible_faqs}) != len(visible_faqs):
        errors.append('FAQ schema consistency')
    removed_schools = set(record['previousSchools']) - set(source_schools(center))
    for node in schema:
        if node.get('@type') in ('Article', 'WebPage', 'Service') and any(x.get('name') in removed_schools for x in node.get('mentions', []) if isinstance(x, dict)):
            errors.append('old mixed-level school mentions')
    for old, new in ((record['previousAddress'], center['address']), (record['previousName'], center['displayName'])):
        if old != new and old not in new and old in raw:
            errors.append('stale exact name/address')
    ids = doc.xpath('//*[@id]/@id')
    if len(ids) != len(set(ids)):
        errors.append('duplicate IDs')
    if align_guide(raw, record, center) != raw:
        errors.append('not idempotent')
    return errors


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument('--write', action='store_true')
    args = parser.parse_args()
    records = targets()
    if len(records) != 106:
        raise ValueError('승인된 106개 페이지 범위가 아닙니다.')
    centers_list = json.loads((ROOT / 'tools/data/branch-directory/branches.json').read_text(encoding='utf-8'))['centers']
    centers = {c['routeName']: c for c in centers_list}
    if len(centers) != len(centers_list):
        raise ValueError('중복 센터명')
    evidence = verify_centers(records, centers)
    stage, backup = REPORT_ROOT / 'staged-site', REPORT_ROOT / 'before-step6'
    destinations = {(ROOT / row['path'].strip('/') / 'index.html').resolve() for row in records}
    protected = [p for p in ROOT.rglob('*.html') if p.resolve() not in destinations]
    protected += list((ROOT / 'assets').rglob('*.*')) + list((ROOT / 'tools/data').rglob('*.json'))
    protected += list(ROOT.glob('*.xml')) + [ROOT / 'robots.txt', SOURCE_WORKBOOK]
    source_dir = Path(r'C:\Users\1992k\Desktop\프로그램 원고')
    protected += [source_dir / f'{level} {subject}학원.zip' for level in ('초등', '중등', '고등') for subject in ('수학', '영어')]
    protected = sorted(set(protected))
    print(f'원본 {len(evidence)}개 센터 확인, 보호 파일 {len(protected)}개 확인 중', flush=True)
    source_hashes = {str(p): sha(p) for p in protected}
    changes = []
    stats = Counter()
    for record in records:
        relative = Path(record['path'].strip('/')) / 'index.html'
        file = (ROOT / relative).resolve()
        if not file.is_file() or not file.is_relative_to((ROOT / '과목별학원').resolve()):
            raise ValueError(f'기존 과목별 페이지 밖 경로: {relative}')
        center = centers[record['center']]
        old = file.read_text(encoding='utf-8')
        new = align_guide(old, record, center)
        if scope_fingerprint(old, record, center) != scope_fingerprint(new, record, center):
            raise ValueError(f'사실 정리 범위 밖 변경: {relative}')
        before_doc, after_doc = html.fromstring(old), html.fromstring(new)
        if Counter(before_doc.xpath('//a/@href')) != Counter(after_doc.xpath('//a/@href')):
            raise ValueError(f'링크 변경: {relative}')
        errors = factual_errors(new, record, center)
        if errors:
            raise ValueError(f'{relative}: {errors}')
        generated = stage / relative
        generated.parent.mkdir(parents=True, exist_ok=True)
        generated.write_text(new, encoding='utf-8')
        baseline = (backup / relative).read_text(encoding='utf-8') if (backup / relative).exists() else old
        previous_schema = graph(html.fromstring(baseline))
        stats['changedThisRun'] += old != new
        stats['pagesChanged'] += baseline != new
        stats['serviceNodesRemoved'] += sum(n.get('@type') == 'Service' for n in previous_schema) - sum(n.get('@type') == 'Service' for n in graph(after_doc))
        stats['unsupportedHoursNodesRemoved'] += sum('openingHoursSpecification' in n for n in previous_schema)
        stats['reviewedInquiryNotices'] += len(after_doc.xpath('//*[@id="guide-course-note"]'))
        stats['addressTextPages'] += record['previousAddress'] != center['address']
        stats['brandNamePages'] += record['previousName'] != center['displayName']
        changes.append({'path': record['path'], 'center': record['center'], 'subject': record['subject'], 'initialDifferences': record['initialDifferences'], 'publishedGrades': course_guidance(center, record['subject'], '고')['grades'], 'pendingGrades': course_guidance(center, record['subject'], '고')['pendingGrades'], 'answer': course_answer(center, record['subject'], '고'), 'beforeSha256': sha(file), 'afterSha256': sha(generated)})
    print(f'{len(changes)}개 페이지의 내용·스키마·멱등성 검증 완료, 보호 파일 재확인 중', flush=True)
    if any(sha(Path(p)) != digest for p, digest in source_hashes.items()):
        raise ValueError('원본 또는 범위 밖 파일 변경')
    for record in changes:
        if sha(ROOT / record['path'].strip('/') / 'index.html') != record['beforeSha256']:
            raise ValueError('작업 중 대상 페이지가 변경되었습니다.')
    if args.write:
        for record in changes:
            relative = Path(record['path'].strip('/')) / 'index.html'
            destination, original = ROOT / relative, backup / relative
            original.parent.mkdir(parents=True, exist_ok=True)
            if not original.exists():
                original.write_bytes(destination.read_bytes())
            if sha(destination) != record['afterSha256']:
                destination.write_bytes((stage / relative).read_bytes())
        if any(sha(Path(p)) != digest for p, digest in source_hashes.items()):
            raise ValueError('적용 후 보호 파일 불일치')
        if any(sha(ROOT / r['path'].strip('/') / 'index.html') != r['afterSha256'] for r in changes):
            raise ValueError('적용 후 대상 파일 불일치')
    report = {'pages': len(records), 'centers': len(evidence), 'counts': dict(stats), 'scopeErrors': 0, 'factualErrors': 0, 'linkChanges': 0, 'protectedFileCount': len(protected), 'nonTargetHtmlCount': sum(p.suffix == '.html' for p in protected), 'sourceEvidence': evidence, 'protectedHashes': source_hashes, 'appliedLocally': args.write, 'deployed': False}
    REPORT_ROOT.mkdir(parents=True, exist_ok=True)
    (REPORT_ROOT / 'verification.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    (REPORT_ROOT / 'changes.json').write_text(json.dumps(changes, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({k: v for k, v in report.items() if k not in ('sourceEvidence', 'protectedHashes')}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
