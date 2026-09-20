"""Generate the reviewed Korean branch directory for wawa-center.kr.

Source workbooks and photos stay immutable. The script writes only the new
``/지점안내/`` tree, its optimized media, a reviewable data manifest, an audit
report, and an isolated sitemap block.

This generator intentionally excludes W+, 글로리드, 대구역점2호관 records
that the site owner previously removed from the public branch directory.
"""

from __future__ import annotations

import hashlib
import json
import re
import shutil
from collections import Counter, defaultdict
from datetime import date
from html import escape
from pathlib import Path
from urllib.parse import quote

from openpyxl import load_workbook
from PIL import Image, ImageOps, ImageStat

from branch_course_guidance import (
    REGISTRATION_NOTE, branch_course_answer, center_notes, course_guidance,
    verify_source, weekend_guidance,
)
from branch_hub_upgrade import upgrade_center, upgrade_directory
from branch_seo import finalize_page, page_dates, save_dates, checked_source_date


ROOT = Path(__file__).resolve().parents[1]
SOURCE_ROOT = Path(r"C:\Users\1992k\Desktop\센터정보")
SOURCE_WORKBOOK = SOURCE_ROOT / "코칭센터_데이터_.xlsx"
TARGET_WORKBOOK = SOURCE_ROOT / "타깃학교 v2.xlsx"
PHOTO_ROOT = SOURCE_ROOT / "센터별 사진"
COMMON_PHOTO_ROOT = PHOTO_ROOT / "1 공용사진"
REFERENCE_IMAGE_ROOT = Path(r"C:\Users\1992k\Desktop\홈페이지 정리\참고자료\공통자료\이미지")
REPRESENTATIVE_ROOT = REFERENCE_IMAGE_ROOT / "대표이미지"
BODY_IMAGE_ROOT = REFERENCE_IMAGE_ROOT / "본문이미지"
MAP_IMAGE_ROOT = REFERENCE_IMAGE_ROOT / "지도이미지"
VERIFIED_MEDIA_MANIFEST = Path(
    r"C:\Users\1992k\Desktop\홈페이지 정리\새 홈페이지15\tools\data\branches\branch-media.json"
)

OUTPUT_ROOT = ROOT / "지점안내"
MEDIA_ROOT = ROOT / "assets" / "branch-directory"
DATA_ROOT = ROOT / "tools" / "data" / "branch-directory"
REPORT_ROOT = ROOT / "reports" / "branch-directory"
SUPPLEMENTAL_CENTERS_FILE = DATA_ROOT / "supplemental-centers.json"

DOMAIN = "https://wawa-center.kr"
PHONE_DISPLAY = "010-3957-8283"
PHONE_LINK = "01039578283"
CONSULT_URL = "https://docs.google.com/forms/d/e/1FAIpQLSdb2oE5Qk5YS0TfYDxyV1w-IOTkhkjOCmmpAKTI9FmqpVj6Yg/viewform"
TODAY = date.today().isoformat()

REGIONS = [
    "서울", "경기", "인천", "부산", "대구", "광주", "대전", "울산",
    "세종", "강원", "충북", "충남", "전북", "전남", "경북", "경남", "제주",
]

REGION_PREFIXES = {
    "서울": ("서울", "서울특별시"),
    "경기": ("경기", "경기도"),
    "인천": ("인천", "인천광역시"),
    "부산": ("부산", "부산광역시"),
    "대구": ("대구", "대구광역시"),
    "광주": ("광주", "광주광역시"),
    "대전": ("대전", "대전광역시"),
    "울산": ("울산", "울산광역시"),
    "세종": ("세종", "세종특별자치시"),
    "강원": ("강원", "강원도", "강원특별자치도"),
    "충북": ("충북", "충청북도"),
    "충남": ("충남", "충청남도"),
    "전북": ("전북", "전라북도", "전북특별자치도"),
    "전남": ("전남", "전라남도"),
    "경북": ("경북", "경상북도"),
    "경남": ("경남", "경상남도"),
    "제주": ("제주", "제주특별자치도"),
}

EXCLUDED_NAMES = {
    "다산점(W+)", "둔산점(W+)", "송도점(W+)", "수지점(W+)",
    "은평점(W+)", "중동점(W+)", "화정점(W+)", "후곡점(W+)",
    "수지점(글로리드)", "은평점(글로리드)", "대구역점2호관",
}

SUBJECTS = ["국어", "영어", "수학", "과학", "사회"]
LEVEL_LABELS = [("초", "초등"), ("중", "중등"), ("고", "고등")]
BRANCH_TOPIC_VARIANTS = [
    ("초등", "수학"), ("초등", "영어"),
    ("중등", "수학"), ("중등", "영어"),
    ("고등", "수학"), ("고등", "영어"),
]
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".gif", ".webp"}
IMAGE_REJECT_TOKENS = {
    "지도", "map", "특강", "홍보", "배너", "banner", "수업료", "교습비",
    "가격", "price", "로고", "logo", "명함", "프로필", "profile", "thumbnail",
}


def clean(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def esc(value: object) -> str:
    return escape(str(value), quote=True)


def clean_html(value: str) -> str:
    """Keep generated markup stable and free of trailing whitespace."""
    return "\n".join(line.rstrip() for line in value.splitlines()) + "\n"


def split_values(value: object) -> list[str]:
    text = clean(value)
    if not text:
        return []
    result: list[str] = []
    for item in re.split(r"\s*[,/·]\s*", text):
        item = item.strip(" -")
        if item and item not in result:
            result.append(item)
    return result


def school_key(value: str) -> str:
    """Collapse common short/full school-name variants within one center."""
    text = clean(value).replace(" ", "")
    text = re.sub(r"^(?:서울|부산|대구|인천|광주|대전|울산|세종)", "", text)
    replacements = [
        ("여자고등학교", "여고"), ("여자중학교", "여중"),
        ("초등학교", "초"), ("중학교", "중"), ("고등학교", "고"),
    ]
    for source, target in replacements:
        text = text.replace(source, target)
    return text


def branch_key(value: object) -> str:
    text = clean(value)
    text = re.sub(
        r"^(?:와와학습코칭센터|와와학습코칭학원|모두오름학습코칭학원|모두오름학원코칭학원)\s*",
        "",
        text,
    )
    return text.replace("(모두)", "").strip()


def route_name(source_name: str) -> str:
    return source_name.replace("(모두)", "").strip()


def encoded_url(path: str) -> str:
    return DOMAIN + quote(path, safe="/#")


def region_from_address(address: str) -> str:
    for region, prefixes in REGION_PREFIXES.items():
        if any(address.startswith(prefix) for prefix in prefixes):
            return region
    raise ValueError(f"주소에서 시·도를 확인할 수 없습니다: {address}")


def district_from_address(address: str, region: str) -> str:
    tokens = address.split()
    if len(tokens) < 2:
        return region
    if region == "세종":
        return "세종시"
    return tokens[1]


def reader_location(value: object) -> str:
    text = clean(value)
    if not text:
        return ""
    text = re.sub(r"^안녕하세요[,! ]*", "", text)
    text = re.sub(r"OO학생\s*학부모님[^~.!?]*[~.!?]*", "", text)
    text = text.replace("부탁드립니다", "확인해 주세요")
    text = re.sub(r"\s+", " ", text).strip(" .")
    if len(text) > 220:
        text = text[:217].rstrip() + "…"
    return text


def weekend_summary(value: object, detail: object) -> str:
    status = clean(value)
    schedule = clean(detail)
    normalized = status.replace(" ", "")
    if not status:
        return "주말 수업 여부는 센터 상담에서 확인해 주세요."
    # A negative for Sunday must not erase an explicitly positive Saturday.
    if "토요일" in status and "일요일" in status and "과학수업만가능" in normalized:
        return "토요일은 과학 수업만 운영하며 일요일 수업은 운영하지 않습니다. " + ("안내 시간: " + schedule.rstrip(".") + "." if schedule else "과목별 시간은 상담에서 확인해 주세요.")
    if "불가" in normalized or "둘다불가" in normalized:
        return "주말 정규 수업은 운영하지 않습니다."
    if "시험대비" in status:
        return status.rstrip(".") + "."
    status = status.replace("가능함", "가능").replace("주말수업가능", "주말 수업 가능")
    status = re.sub(r"토요일\s*가능", "토요일 가능", status)
    status = re.sub(r"일요일\s*가능", "일요일 가능", status)
    result = status.rstrip(".") + "."
    if schedule and schedule not in {"협의", "협의 부탁드립니다.", "불가", "주말불가"}:
        result += " 참고 시간: " + schedule.rstrip(".") + "."
    elif any(token in normalized for token in ("가능", "조율", "협의", "토요일", "일요일")):
        result += " 과목별 시간은 상담에서 확인해 주세요."
    return result


def subject_level_summary(subjects: dict[str, list[str]]) -> str:
    parts = []
    for subject in SUBJECTS:
        grades = subjects.get(subject, [])
        if not grades:
            continue
        levels = [label for prefix, label in LEVEL_LABELS if any(g.startswith(prefix) for g in grades)]
        if levels:
            parts.append(f"{subject} {'·'.join(levels)}")
    return ", ".join(parts) or "가능 과목과 학년은 상담 시 확인"


def safe_management_points(value: object) -> list[tuple[str, str]]:
    text = clean(value)
    rules = [
        (("플래너", "계획표"), "플래너 점검", "학습 계획과 실행 결과를 함께 확인하는 운영 항목이 기록돼 있습니다."),
        (("오답",), "오답 재학습", "틀린 이유를 구분하고 다시 풀어보는 관리 항목을 상담에서 확인할 수 있습니다."),
        (("내신", "학교별"), "학교별 내신 준비", "재학 학교와 시험 범위를 기준으로 준비 방향을 확인하는 항목이 있습니다."),
        (("시험", "기출"), "시험기간 계획", "시험 범위와 일정에 맞춘 학습 계획 운영 여부를 확인할 수 있습니다."),
        (("학습일지", "평가서"), "학습 기록", "학습일지 또는 평가 기록을 활용하는 운영 항목이 기재돼 있습니다."),
        (("피드백", "소통", "단체카톡방", "단톡방"), "보호자 피드백", "출결과 학습 진행 상황을 공유하는 방식은 상담에서 구체적으로 확인해 주세요."),
        (("독서",), "독서 활동", "독서 활동을 학습 과정에 연결하는 운영 항목이 기재돼 있습니다."),
        (("입시",), "고등 학습 상담", "고등 학습 계획과 진로·입시 상담 범위는 센터에서 확인할 수 있습니다."),
        (("자습",), "자습 운영", "시험기간 자습 운영 여부와 이용 조건은 센터 상담에서 확인해 주세요."),
        (("AI", "ai"), "학습 도구 활용", "AI 학습 도구의 적용 과목과 이용 방식은 센터에서 확인할 수 있습니다."),
        (("전과목", "국영수사과"), "과목 연계 관리", "여러 과목의 학습량과 일정 조율 범위를 상담에서 확인할 수 있습니다."),
    ]
    result = []
    for tokens, title, copy in rules:
        if any(token in text for token in tokens):
            result.append((title, copy))
    return result[:4]


def load_target_rows() -> dict[str, dict[str, object]]:
    wb = load_workbook(TARGET_WORKBOOK, read_only=True, data_only=True)
    ws = wb.active
    grouped: dict[str, dict[str, object]] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0] or not row[3]:
            continue
        key = branch_key(row[3])
        item = grouped.setdefault(
            key,
            {"neighborhoods": [], "regions": [], "districts": [], "schools": {"초등": [], "중등": [], "고등": []}},
        )
        for name, value in (("neighborhoods", row[0]), ("regions", row[1]), ("districts", row[2])):
            value = clean(value)
            if value and value not in item[name]:
                item[name].append(value)
        for level, value in zip(("초등", "중등", "고등"), row[4:7]):
            for school in split_values(value):
                if school not in item["schools"][level]:
                    item["schools"][level].append(school)
    return grouped


def load_centers() -> tuple[list[dict[str, object]], list[dict[str, str]]]:
    verify_source(SOURCE_WORKBOOK)
    target = load_target_rows()
    wb = load_workbook(SOURCE_WORKBOOK, read_only=True, data_only=True)
    ws = wb.active
    centers: list[dict[str, object]] = []
    excluded: list[dict[str, str]] = []

    for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        source_name = clean(row[0])
        if not source_name:
            continue
        if source_name in EXCLUDED_NAMES:
            excluded.append({"sourceRow": str(row_number), "name": source_name, "reason": "기존 공개 지점안내 삭제 대상"})
            continue

        key = branch_key(source_name)
        mapped = target.get(key, {})
        address = clean(row[11])
        region = clean((mapped.get("regions") or [""])[0]) or region_from_address(address)
        if region not in REGIONS:
            region = region_from_address(address)
        district = clean((mapped.get("districts") or [""])[0]) or district_from_address(address, region)
        name = route_name(source_name)
        is_modu = "(모두)" in source_name or "모두오름" in clean(row[6])
        display_name = f"모두오름학습코칭학원 {name}" if is_modu else f"와와학습코칭센터 {name}"

        subjects: dict[str, list[str]] = {}
        for subject, value in zip(SUBJECTS, row[16:21]):
            grades = split_values(value)
            if grades:
                subjects[subject] = grades

        workbook_schools = {
            level: split_values(value)
            for level, value in zip(("초등", "중등", "고등"), row[13:16])
        }
        schools = {"초등": [], "중등": [], "고등": []}
        for level in schools:
            seen_school_keys = set()
            for school in (mapped.get("schools", {}).get(level, []) if mapped else []) + workbook_schools[level]:
                school = clean(school)
                key_name = school_key(school)
                if school and key_name not in seen_school_keys:
                    schools[level].append(school)
                    seen_school_keys.add(key_name)

        center = {
            "sourceRow": row_number,
            "sourceName": source_name,
            "routeName": name,
            "key": key,
            "brand": "모두오름학습코칭학원" if is_modu else "와와학습코칭센터",
            "displayName": display_name,
            "region": region,
            "district": district,
            "address": address,
            "locationGuide": reader_location(row[12]),
            "registeredName": clean(row[6]),
            "registrationNumber": clean(row[7]),
            "registrationDate": clean(row[8]),
            "subjects": subjects,
            "schools": schools,
            "neighborhoods": list(mapped.get("neighborhoods", [])) if mapped else [],
            "openingReference": clean(row[26]),
            "weekend": weekend_summary(row[27], row[28]),
            "managementPoints": safe_management_points(row[29]),
            "photoSource": str(PHOTO_ROOT / source_name),
            "photos": [],
        }
        centers.append(center)

    # A separately reviewed source can fill a workbook omission without
    # changing the original workbook or borrowing a different branch's row.
    supplemental = json.loads(SUPPLEMENTAL_CENTERS_FILE.read_text(encoding="utf-8"))["centers"]
    for record in supplemental:
        center = dict(record)
        name = center["routeName"]
        key = branch_key(name)
        if name in EXCLUDED_NAMES or any(item["routeName"] == name for item in centers):
            raise ValueError(f"추가 센터가 기존 또는 제외 센터와 중복됩니다: {name}")
        mapped = target.get(key)
        if not mapped or not mapped["neighborhoods"]:
            raise ValueError(f"추가 센터의 타깃 동네가 없습니다: {name}")
        if center["region"] != region_from_address(center["address"]):
            raise ValueError(f"추가 센터의 주소와 지역이 다릅니다: {name}")
        if not center.get("verifiedMediaKey") or not center.get("sourceProvenance"):
            raise ValueError(f"추가 센터의 검증 출처가 없습니다: {name}")
        center.update({
            "sourceRow": None,
            "sourceName": name,
            "key": key,
            "displayName": f'{center["brand"]} {name}',
            "neighborhoods": list(mapped["neighborhoods"]),
            "schools": mapped["schools"],
            "photoSource": str(PHOTO_ROOT / name),
            "photos": [],
        })
        centers.append(center)

    centers.sort(key=lambda item: (REGIONS.index(item["region"]), item["routeName"]))
    return centers, excluded


def image_score(path: Path) -> tuple[float, int, int] | None:
    lowered = path.name.lower()
    if any(token.lower() in lowered for token in IMAGE_REJECT_TOKENS):
        return None
    try:
        with Image.open(path) as image:
            image = ImageOps.exif_transpose(image).convert("RGB")
            width, height = image.size
            sample = image.copy()
            sample.thumbnail((128, 128))
            hsv = sample.convert("HSV")
            saturation = ImageStat.Stat(hsv.getchannel("S")).mean[0] / 255
            pixels = list(hsv.getdata())
            pale_fraction = sum(1 for _, sat, val in pixels if sat < 24 and val > 205) / max(len(pixels), 1)
    except Exception:
        return None
    if width < 480 or height < 320:
        return None
    ratio = width / max(height, 1)
    if ratio < 0.62 or ratio > 2.4:
        return None
    megapixels = min((width * height) / 1_000_000, 8)
    orientation = 2.2 - abs(ratio - 1.45)
    name_bonus = 0.4 if any(token in lowered for token in ("사진", "센터", "photo", "kakao")) else 0
    visual_score = saturation * 3.2 - pale_fraction * 2.2
    return (megapixels + orientation + name_bonus + visual_score, width, height)


def select_photos(directory: Path, limit: int = 4) -> list[Path]:
    if not directory.exists():
        return []
    candidates = []
    seen_hashes = set()
    for path in directory.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in IMAGE_EXTENSIONS:
            continue
        score = image_score(path)
        if not score:
            continue
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        if digest in seen_hashes:
            continue
        seen_hashes.add(digest)
        candidates.append((score, path))
    candidates.sort(key=lambda item: (-item[0][0], item[1].name.lower()))
    return [path for _, path in candidates[:limit]]


def optimize_photo(source: Path, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(source) as original:
        image = ImageOps.exif_transpose(original)
        if image.mode not in {"RGB", "RGBA"}:
            image = image.convert("RGB")
        if image.mode == "RGBA":
            background = Image.new("RGB", image.size, "white")
            background.paste(image, mask=image.getchannel("A"))
            image = background
        if image.width > 1440:
            new_height = round(image.height * 1440 / image.width)
            image = image.resize((1440, new_height), Image.Resampling.LANCZOS)
        image.save(destination, "WEBP", quality=82, method=6)


def media_dimensions(path: Path) -> tuple[int, int]:
    with Image.open(path) as image:
        return image.size


def copy_primary_asset(source: Path, folder: str) -> dict[str, object]:
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    extension = source.suffix.lower()
    destination = MEDIA_ROOT / folder / f"{digest[:16]}{extension}"
    destination.parent.mkdir(parents=True, exist_ok=True)
    if not destination.exists():
        shutil.copy2(source, destination)
    width, height = media_dimensions(source)
    return {
        "src": "/" + destination.relative_to(ROOT).as_posix(),
        "width": width,
        "height": height,
        "sha256": digest,
    }


def responsive_body_asset(source: Path) -> dict[str, object]:
    """Copy the original body image and create full-height mobile variants.

    The artwork remains uncropped. AVIF and WebP variants only reduce transfer
    size; the original JPEG remains the fallback for older browsers.
    """
    asset = copy_primary_asset(source, "body")
    digest = asset["sha256"]
    variants: dict[str, list[dict[str, object]]] = {"avif": [], "webp": []}
    with Image.open(source) as original:
        image = ImageOps.exif_transpose(original).convert("RGB")
        widths = sorted({width for width in (480, 768, image.width) if width <= image.width})
        for width in widths:
            if width == image.width:
                resized = image
            else:
                height = round(image.height * width / image.width)
                resized = image.resize((width, height), Image.Resampling.LANCZOS)
            for extension, image_format, options in (
                ("avif", "AVIF", {"quality": 52, "speed": 6}),
                ("webp", "WEBP", {"quality": 80, "method": 6}),
            ):
                destination = MEDIA_ROOT / "body" / f'{digest[:16]}-{width}.{extension}'
                destination.parent.mkdir(parents=True, exist_ok=True)
                if not destination.exists():
                    resized.save(destination, image_format, **options)
                variants[extension].append({
                    "src": "/" + destination.relative_to(ROOT).as_posix(),
                    "width": width,
                })
    asset["variants"] = variants
    return asset


def import_primary_media(centers: list[dict[str, object]]) -> None:
    representative_sources = sorted(
        (
            path for path in REPRESENTATIVE_ROOT.rglob("*")
            if path.is_file() and path.suffix.lower() in IMAGE_EXTENSIONS
        ),
        key=lambda path: (path.name.casefold(), path.as_posix().casefold()),
    )
    if not representative_sources:
        raise RuntimeError(f"대표이미지를 찾을 수 없습니다: {REPRESENTATIVE_ROOT}")

    verified = json.loads(VERIFIED_MEDIA_MANIFEST.read_text(encoding="utf-8"))["centers"]
    body_cache: dict[str, dict[str, object]] = {}
    representative_cache: dict[Path, dict[str, object]] = {}
    map_cache: dict[Path, dict[str, object]] = {}

    for center in centers:
        seed = int(
            hashlib.sha256(
                f'{center["region"]}/{center["routeName"]}'.encode("utf-8")
            ).hexdigest()[:16],
            16,
        )
        representative_source = representative_sources[seed % len(representative_sources)]
        if representative_source not in representative_cache:
            representative_cache[representative_source] = copy_primary_asset(
                representative_source, "representative"
            )
        representative = representative_cache[representative_source]

        body_name = "seoul3957.jpg" if center["region"] == "서울" else "local3957.jpg"
        body_source = BODY_IMAGE_ROOT / body_name
        if not body_source.is_file():
            raise FileNotFoundError(body_source)
        if body_name not in body_cache:
            body_cache[body_name] = responsive_body_asset(body_source)
        body = body_cache[body_name]

        media_key = center.get("verifiedMediaKey") or f'center-row-{int(center["sourceRow"]):03d}'
        map_spec = verified.get(media_key, {}).get("map")
        if not map_spec:
            raise RuntimeError(f'검증 지도 매칭이 없습니다: {center["routeName"]} ({media_key})')
        if center.get("verifiedMediaKey"):
            match = verified[media_key].get("mapMatch", {})
            registration = re.sub(r"\D", "", center["registrationNumber"])
            if (match.get("reviewStatus") != "confirmed"
                    or match.get("centerName") != center["routeName"]
                    or re.sub(r"\D", "", match.get("centerRegistrationNumber", "")) != registration):
                raise RuntimeError(f'추가 센터 지도 검증 정보가 다릅니다: {center["routeName"]}')
        map_source = MAP_IMAGE_ROOT / Path(map_spec["src"]).name
        if not map_source.is_file():
            map_source = Path(map_spec["sourcePath"])
        if not map_source.is_file():
            raise FileNotFoundError(f'지도 파일을 찾을 수 없습니다: {center["routeName"]} / {map_source}')
        if center.get("verifiedMediaKey") and hashlib.sha256(map_source.read_bytes()).hexdigest() != map_spec["sha256"]:
            raise RuntimeError(f'추가 센터 지도 원본이 변경되었습니다: {center["routeName"]}')
        if map_source not in map_cache:
            map_cache[map_source] = copy_primary_asset(map_source, "maps")
        map_asset = map_cache[map_source]

        center["primaryMedia"] = {
            "representative": representative,
            "body": body,
            "map": map_asset,
            "mapMatch": "verified-reference-center" if center.get("verifiedMediaKey") else "verified-center-source-row",
        }
        center["informationReviewedAt"] = checked_source_date(center)


def import_media(centers: list[dict[str, object]]) -> None:
    if MEDIA_ROOT.exists():
        shutil.rmtree(MEDIA_ROOT)

    common_sources = select_photos(COMMON_PHOTO_ROOT, limit=4)
    if not common_sources:
        raise RuntimeError(f"공용 센터 사진을 찾을 수 없습니다: {COMMON_PHOTO_ROOT}")
    common_hashes = {hashlib.sha256(path.read_bytes()).hexdigest() for path in common_sources}
    common_urls = []
    for index, source in enumerate(common_sources, start=1):
        destination = MEDIA_ROOT / "common" / f"center-{index:02d}.webp"
        optimize_photo(source, destination)
        common_urls.append("/" + destination.relative_to(ROOT).as_posix())

    for center in centers:
        source_dir = Path(center["photoSource"])
        selected = select_photos(source_dir)
        selected_hashes = {hashlib.sha256(path.read_bytes()).hexdigest() for path in selected}
        uses_only_common_copies = bool(selected) and selected_hashes.issubset(common_hashes)

        if not selected or uses_only_common_copies:
            offset = int(hashlib.sha256(center["key"].encode("utf-8")).hexdigest()[:8], 16) % len(common_urls)
            center["photos"] = common_urls[offset:] + common_urls[:offset]
            center["photoMode"] = "common"
            continue

        media_dir = MEDIA_ROOT / center["region"] / center["routeName"]
        urls = []
        for index, source in enumerate(selected, start=1):
            destination = media_dir / f"photo-{index:02d}.webp"
            optimize_photo(source, destination)
            urls.append("/" + destination.relative_to(ROOT).as_posix())
        center["photos"] = urls
        center["photoMode"] = "center"


def breadcrumb(items: list[tuple[str, str]]) -> str:
    parts = []
    for index, (label, href) in enumerate(items):
        if index == len(items) - 1:
            parts.append(f'<li><span aria-current="page">{esc(label)}</span></li>')
        else:
            parts.append(f'<li><a href="{esc(href)}">{esc(label)}</a></li>')
    return '<nav class="branch-breadcrumb" aria-label="현재 위치"><ol>' + "".join(parts) + "</ol></nav>"


def site_header(active: str = "branches") -> str:
    links = [
        ("홈", "/"), ("학원소개", "/overview/"), ("학습가이드", "/guide/"),
        ("교육정보", "/교육정보/"), ("학부모후기", "/학부모후기/"),
        ("과목별학원", "/과목별학원/"), ("학년별학원", "/학년별학원/"),
        ("지점안내", "/지점안내/"),
    ]
    nav = "".join(
        f'<a{" class=\"active\"" if active == "branches" and label == "지점안내" else ""} href="{href}">{label}</a>'
        for label, href in links
    )
    return (
        '<header class="site-header"><nav class="nav" aria-label="주요 메뉴">'
        '<a class="logo" href="/"><span class="brand-orange">와와</span>학습'
        '<span class="brand-orange">코칭</span>센터 <span class="brand-tail">영어수학 전문학원</span></a>'
        f'<div class="nav-links" aria-label="페이지 이동">{nav}</div></nav></header>'
    )


def site_footer() -> str:
    return f'''<footer class="branch-footer">
  <strong>와와학습코칭센터</strong>
  <p>센터별 운영 과목·학년·요일은 다를 수 있으므로 방문 전 상담에서 확인해 주세요.</p>
  <a href="tel:{PHONE_LINK}">{PHONE_DISPLAY}</a>
</footer>
<div class="wawa-fixed-fab-container" aria-label="빠른 상담">
  <a href="tel:{PHONE_DISPLAY}" class="wawa-fab-item fab-call"><span class="fab-icon" aria-hidden="true">☎</span><span class="fab-text">전화문의</span></a>
  <a href="https://blogsms.net/{PHONE_LINK}" target="_blank" rel="noopener noreferrer" class="wawa-fab-item fab-sms"><span class="fab-icon" aria-hidden="true">✉</span><span class="fab-text">문자문의</span></a>
  <a href="{CONSULT_URL}" target="_blank" rel="noopener noreferrer" class="wawa-fab-item fab-consult"><span class="fab-icon" aria-hidden="true">✓</span><span class="fab-text">상담신청</span></a>
</div>'''


def page_head(title: str, description: str, canonical_path: str, graph: list[dict[str, object]], image: str) -> str:
    canonical = encoded_url(canonical_path)
    image_url = image if image.startswith("http") else DOMAIN + quote(image, safe="/")
    json_ld = json.dumps({"@context": "https://schema.org", "@graph": graph}, ensure_ascii=False, separators=(",", ":"))
    return f'''<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{esc(title)}</title>
  <meta name="description" content="{esc(description)}">
  <meta name="robots" content="index, follow, max-image-preview:large">
  <link rel="canonical" href="{canonical}">
  <link rel="icon" type="image/png" href="/assets/favicon.png">
  <link rel="apple-touch-icon" href="/assets/favicon.png">
  <link rel="alternate" type="application/rss+xml" href="{DOMAIN}/rss.xml" title="와와학습코칭센터 RSS">
  <meta property="og:locale" content="ko_KR">
  <meta property="og:type" content="website">
  <meta property="og:site_name" content="와와학습코칭센터">
  <meta property="og:title" content="{esc(title)}">
  <meta property="og:description" content="{esc(description)}">
  <meta property="og:url" content="{canonical}">
  <meta property="og:image" content="{image_url}">
  <meta name="twitter:card" content="summary_large_image">
  <link rel="stylesheet" href="/assets/header.css">
  <link rel="stylesheet" href="/assets/fab.css">
  <link rel="stylesheet" href="/assets/branch-directory.css?v=20260920d">
  <script type="application/ld+json">{json_ld}</script>
</head>'''


def page_shell(title: str, description: str, path: str, graph: list[dict[str, object]], image: str, crumbs: list[tuple[str, str]], body: str) -> str:
    return f'''<!doctype html>
<html lang="ko">
{page_head(title, description, path, graph, image)}
<body class="branch-directory-page">
{site_header()}
{breadcrumb(crumbs)}
<main id="main" class="branch-shell">{body}</main>
{site_footer()}
<script src="/assets/branch-directory.js?v=20260920" defer></script>
</body>
</html>
'''


def graph_base(title: str, description: str, path: str, crumbs: list[tuple[str, str]], page_type: str = "WebPage") -> list[dict[str, object]]:
    page_url = encoded_url(path)
    return [
        {"@type": "WebSite", "@id": DOMAIN + "/#website", "url": DOMAIN + "/", "name": "와와학습코칭센터", "inLanguage": "ko-KR", "publisher": {"@id": DOMAIN + "/#organization"}},
        {"@type": "EducationalOrganization", "@id": DOMAIN + "/#organization", "url": DOMAIN + "/", "name": "와와학습코칭센터", "logo": DOMAIN + "/assets/favicon.png", "telephone": "+82-10-3957-8283", "areaServed": {"@type": "Country", "name": "대한민국"}},
        {
            "@type": page_type,
            "@id": page_url + "#webpage",
            "url": page_url,
            "name": title,
            "description": description,
            "inLanguage": "ko-KR",
            "isPartOf": {"@id": DOMAIN + "/#website"},
            "breadcrumb": {"@id": page_url + "#breadcrumb"},
        },
        {
            "@type": "BreadcrumbList",
            "@id": page_url + "#breadcrumb",
            "itemListElement": [
                {"@type": "ListItem", "position": index + 1, "name": label, "item": encoded_url(href)}
                for index, (label, href) in enumerate(crumbs)
            ],
        },
    ]


def branch_card(center: dict[str, object]) -> str:
    path = f'/지점안내/{center["region"]}/{center["routeName"]}/'
    neighborhoods = center["neighborhoods"][:4]
    search = " ".join([center["displayName"], center["region"], center["district"], *center["neighborhoods"]]).lower()
    return f'''<article class="branch-card" data-branch-card data-search="{esc(search)}">
  <a href="{esc(path)}">
    <span class="branch-card-region">{esc(center["region"])} · {esc(center["district"])}</span>
    <h3>{esc(center["displayName"])}</h3>
    <p>{esc(center["address"])}</p>
    {f'<small>{esc(" · ".join(neighborhoods))}</small>' if neighborhoods else ''}
    <b>센터 정보 보기</b>
  </a>
</article>'''


def search_panel(label: str) -> str:
    return f'''<section class="branch-search-panel" aria-labelledby="branch-search-title">
  <div>
    <p class="branch-kicker">CENTER SEARCH</p>
    <h2 id="branch-search-title">{esc(label)} 센터 검색</h2>
    <p>지점명, 시·군·구, 수업 가능 동네를 입력해 가까운 센터를 찾아보세요.</p>
  </div>
  <label for="branch-search-input">센터 검색어</label>
  <div class="branch-search-field">
    <input id="branch-search-input" type="search" data-branch-search-input placeholder="예: 명일점, 강동구, 천호동" autocomplete="off">
    <button type="button" data-branch-search-reset disabled>초기화</button>
  </div>
  <p class="branch-search-status" data-branch-search-status role="status" aria-live="polite"></p>
</section>'''


def coaching_system_section(scope: str, intro: str) -> str:
    return f'''<section class="branch-section branch-system" id="coaching-system" aria-labelledby="coaching-system-title">
  <div class="branch-section-head"><p class="branch-kicker">LEARNING SYSTEM</p><h2 id="coaching-system-title">{esc(scope)}에서 확인할 학습코칭 방식</h2><p>{esc(intro)}</p></div>
  <div class="branch-system-grid">
    <article><span>PLAN</span><h3>플랜관리</h3><p>학생과 학습 목표·우선순위를 정하고, 공부 시간과 분량을 스스로 실행할 수 있도록 계획을 구체화합니다.</p></article>
    <article><span>STUDY</span><h3>학습관리</h3><p>현재 이해도와 진도에 맞춰 교재와 학습법을 조정하고, 오답노트·백지노트·마인드맵 등 필요한 복습 도구를 연결합니다.</p></article>
    <article><span>ROUTINE</span><h3>생활관리</h3><p>수업 밖의 학습 습관과 생활 리듬도 함께 살피며, 학생·보호자와의 소통을 다음 계획에 반영합니다.</p></article>
  </div>
  <div class="branch-ai-summary"><div><p class="branch-kicker">AI LEARNING</p><h3>진단에서 취약점 훈련까지</h3><p>공식 학습 시스템은 영어·수학·국어·독서 영역에서 진단, 성취도·취약점 분석, 맞춤 훈련으로 이어지는 학습 도구를 안내합니다. 과목·학년과 실제 활용 여부는 지점별로 다를 수 있으므로 각 센터 페이지와 상담에서 확인해 주세요.</p></div><ul><li>영어: 레벨과 영역별 학습 점검</li><li>수학: 성취도 기반 문제와 오답 클리닉</li><li>국어: 취약 유형 분석과 맞춤 문제</li><li>독서: 독서 기록과 습관 관리</li></ul></div>
  <div class="branch-inline-links"><a href="/과목별학원/">과목별 학습 안내</a><a href="/학년별학원/">학년별 학습 안내</a><a href="/guide/">학습가이드</a></div>
</section>'''


def generate_hub(centers: list[dict[str, object]]) -> str:
    path = "/지점안내/"
    title = "전국 지점안내 | 와와학습코칭센터"
    description = f"전국 {len(centers)}개 와와학습코칭센터·모두오름학습코칭학원 지점의 주소, 가능 과목·학년, 인근 학교와 상담 정보를 지역별로 확인하세요."
    crumbs = [("홈", "/"), ("지점안내", path)]
    graph = graph_base(title, description, path, crumbs, "CollectionPage")
    webpage = next(item for item in graph if item.get("@id") == encoded_url(path) + "#webpage")
    webpage["about"] = [
        {"@type": "Thing", "name": "전국 학습코칭센터"},
        {"@type": "Thing", "name": "플랜관리"},
        {"@type": "Thing", "name": "학습관리"},
        {"@type": "Thing", "name": "생활관리"},
    ]
    webpage["mentions"] = [
        {"@type": "Thing", "name": "AI 영어"},
        {"@type": "Thing", "name": "AI 수학"},
        {"@type": "Thing", "name": "AI 국어"},
        {"@type": "Thing", "name": "AI 독서"},
    ]
    webpage["hasPart"] = [
        {"@type": "WebPageElement", "@id": encoded_url(path) + "#coaching-system", "name": "학습코칭 방식"},
        {"@type": "WebPageElement", "@id": encoded_url(path) + "#branch-list", "name": "전체 센터 목록"},
        {"@type": "WebPageElement", "@id": encoded_url(path) + "#faq", "name": "지점 선택 FAQ"},
    ]
    region_groups = {region: [c for c in centers if c["region"] == region] for region in REGIONS}
    region_items = [(region, group) for region, group in region_groups.items() if group]
    graph.append({
        "@type": "ItemList", "@id": encoded_url(path) + "#regions", "name": "지역별 지점안내",
        "numberOfItems": len(region_items),
        "itemListElement": [
            {"@type": "ListItem", "position": index + 1, "name": f"{region} 지점안내", "url": encoded_url(f"/지점안내/{region}/")}
            for index, (region, _) in enumerate(region_items)
        ],
    })
    faqs = [
        ("가까운 센터는 어떻게 찾나요?", "지점명, 시·군·구, 동네명을 검색하면 관련 센터를 확인할 수 있습니다. 센터 페이지에서 정확한 주소와 수업 가능 동네를 다시 확인해 주세요."),
        ("모든 센터의 수업 과목과 학년이 같은가요?", "아닙니다. 센터별 코치 구성과 시간표에 따라 가능 과목과 학년이 다릅니다. 각 센터 페이지의 과목표를 확인한 뒤 상담에서 현재 운영 여부를 확인해 주세요."),
        ("상담 전에 무엇을 준비하면 좋나요?", "학생의 학년과 학교, 최근 시험지, 현재 교재, 어려운 단원, 평소 숙제와 공부 시간을 정리하면 상담 방향을 빠르게 잡는 데 도움이 됩니다."),
        ("AI 학습 도구는 모든 센터에서 동일하게 이용하나요?", "영어·수학·국어·독서의 진단·분석·훈련 도구가 안내되어 있지만 실제 개설 과목과 활용 범위는 센터별 운영 상황에 따라 다를 수 있습니다. 희망 과목과 학년을 지점 상담에서 확인해 주세요."),
    ]
    graph.append({"@type": "FAQPage", "@id": encoded_url(path) + "#faq", "mainEntity": [
        {"@type": "Question", "name": q, "acceptedAnswer": {"@type": "Answer", "text": a}} for q, a in faqs
    ]})

    region_cards = "".join(
        f'<a class="region-card" href="/지점안내/{esc(region)}/"><span>{esc(region)}</span><strong>{len(group)}개 센터</strong><small>지역 지점 보기</small></a>'
        for region, group in region_items
    )
    cards = "".join(branch_card(center) for center in centers)
    faq_html = "".join(f'<details><summary>{esc(q)}</summary><p>{esc(a)}</p></details>' for q, a in faqs)
    body = f'''
<section class="branch-hub-hero">
  <p class="branch-kicker">WAWA BRANCH DIRECTORY</p>
  <h1>전국 와와학습코칭센터 지점안내</h1>
  <p>{esc(description)}</p>
  <div class="branch-hero-actions"><a class="branch-button primary" href="#branch-list">센터 찾기</a><a class="branch-button" href="/guide/consultation-diagnosis/">상담 준비 가이드</a></div>
</section>
<section class="branch-answer" aria-labelledby="hub-answer-title">
  <p class="branch-kicker">QUICK ANSWER</p><h2 id="hub-answer-title">주소와 실제 운영 범위를 기준으로 찾는 지점 허브</h2>
  <p>지역 목록에서 가까운 센터를 선택하면 등록 명칭·주소·가능 과목과 학년·인근 학교·주말 운영 참고사항을 한 페이지에서 확인할 수 있습니다.</p>
</section>
{coaching_system_section("전국 지점안내", "가까운 센터를 찾는 것과 함께 학생의 계획·학습·생활 관리가 어떤 흐름으로 이어지는지 확인할 수 있도록 공식 학습 시스템의 핵심을 정리했습니다.")}
<section class="branch-section" aria-labelledby="region-list-title"><div class="branch-section-head"><p class="branch-kicker">REGIONS</p><h2 id="region-list-title">시·도별 지점</h2><p>지역을 선택하면 해당 시·도의 센터만 모아 볼 수 있습니다.</p></div><div class="region-grid">{region_cards}</div></section>
{search_panel("전국")}
<section class="branch-section" id="branch-list" aria-labelledby="branch-list-title"><div class="branch-section-head"><p class="branch-kicker">ALL CENTERS</p><h2 id="branch-list-title">전체 센터</h2><p data-branch-count>{len(centers)}개 센터를 표시하고 있습니다.</p></div><div class="branch-card-grid" data-branch-grid>{cards}</div><p class="branch-empty" data-branch-empty hidden>일치하는 센터가 없습니다. 검색어를 줄여 다시 확인해 주세요.</p></section>
<section class="branch-section branch-method" aria-labelledby="method-title"><div class="branch-section-head"><p class="branch-kicker">COACHING FLOW</p><h2 id="method-title">상담에서 학습 계획까지 확인하는 순서</h2></div><ol class="method-grid"><li><b>01</b><strong>현재 상태 확인</strong><span>학교·학년, 최근 시험, 어려운 과목과 단원을 정리합니다.</span></li><li><b>02</b><strong>진도와 학습량 설정</strong><span>교재와 시작 단원, 주간 학습량을 학생 상황에 맞춰 확인합니다.</span></li><li><b>03</b><strong>실행과 오답 점검</strong><span>플래너와 문제풀이 결과를 보고 막힌 원인을 다시 확인합니다.</span></li><li><b>04</b><strong>다음 계획 조정</strong><span>수업과 과제 결과를 바탕으로 다음 학습 계획을 조정합니다.</span></li></ol></section>
<section class="branch-section branch-faq" id="faq" aria-labelledby="hub-faq-title"><div class="branch-section-head"><p class="branch-kicker">FAQ</p><h2 id="hub-faq-title">지점 선택 전 자주 묻는 질문</h2></div>{faq_html}</section>
'''
    html = finalize_page(upgrade_directory(page_shell(title, description, path, graph, "/assets/title.png", crumbs, body), centers), path)
    output = OUTPUT_ROOT / "index.html"
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(clean_html(html), encoding="utf-8")
    return path


def generate_region_page(region: str, centers: list[dict[str, object]]) -> str:
    path = f"/지점안내/{region}/"
    title = f"{region} 지점안내 | 와와학습코칭센터"
    description = f"{region} 지역 {len(centers)}개 학습코칭 센터의 주소, 가능 과목·학년, 수업 가능 동네와 인근 학교를 비교해 가까운 지점을 찾아보세요."
    crumbs = [("홈", "/"), ("지점안내", "/지점안내/"), (region, path)]
    graph = graph_base(title, description, path, crumbs, "CollectionPage")
    webpage = next(item for item in graph if item.get("@id") == encoded_url(path) + "#webpage")
    webpage["about"] = [
        {"@type": "Place", "name": region},
        {"@type": "Thing", "name": f"{region} 학습코칭센터"},
        {"@type": "Thing", "name": "개별 맞춤 학습관리"},
    ]
    webpage["mentions"] = [
        {"@type": "Thing", "name": "플랜관리"},
        {"@type": "Thing", "name": "오답 재학습"},
        {"@type": "Thing", "name": "AI 학습 도구"},
    ]
    webpage["hasPart"] = [
        {"@type": "WebPageElement", "@id": encoded_url(path) + "#coaching-system", "name": f"{region} 학습코칭 방식"},
        {"@type": "WebPageElement", "@id": encoded_url(path) + "#branch-list", "name": f"{region} 센터 목록"},
        {"@type": "WebPageElement", "@id": encoded_url(path) + "#faq", "name": f"{region} 지점 FAQ"},
    ]
    graph.append({
        "@type": "ItemList", "@id": encoded_url(path) + "#centers", "name": f"{region} 센터 목록", "numberOfItems": len(centers),
        "itemListElement": [
            {"@type": "ListItem", "position": index + 1, "name": center["displayName"], "url": encoded_url(f'/지점안내/{region}/{center["routeName"]}/')}
            for index, center in enumerate(centers)
        ],
    })
    districts = Counter(center["district"] for center in centers)
    district_copy = " · ".join(f"{name} {count}곳" for name, count in districts.most_common())
    subject_counts = {
        subject: sum(bool(course_guidance(center, subject)["grades"]) for center in centers)
        for subject in SUBJECTS
    }
    subject_copy = " · ".join(
        f"{subject} {count}곳" for subject, count in subject_counts.items() if count
    )
    region_faqs = [
        (f"{region}에서 가까운 센터는 어떻게 찾나요?", f"{region} 센터 목록에서 지점명, 시·군·구 또는 수업 가능 동네를 검색하세요. 선택한 센터 페이지에서 주소와 위치 안내를 확인한 뒤 방문 전 상담으로 시간표를 확인하면 됩니다."),
        (f"{region} 센터에서 상담할 수 있는 과목은 무엇인가요?", f"센터 자료와 수업 조건을 대조하면 안내 학년이 확인된 곳은 {subject_copy}입니다. 같은 센터가 여러 과목에 포함될 수 있습니다. 현재 시간표와 신규 등록 가능 여부는 상세 페이지와 상담에서 확인해 주세요."),
        ("센터를 비교할 때 무엇을 먼저 보면 좋나요?", "통학 가능한 주소인지 확인한 다음 학생의 학년과 희망 과목, 최근 시험의 반복 오답, 가능한 요일과 시간을 함께 비교하면 상담 범위를 구체화하기 좋습니다."),
    ]
    graph.append({"@type": "FAQPage", "@id": encoded_url(path) + "#faq", "mainEntity": [
        {"@type": "Question", "name": q, "acceptedAnswer": {"@type": "Answer", "text": a}}
        for q, a in region_faqs
    ]})
    cards = "".join(branch_card(center) for center in centers)
    region_faq_html = "".join(f'<details><summary>{esc(q)}</summary><p>{esc(a)}</p></details>' for q, a in region_faqs)
    body = f'''
<section class="branch-hub-hero compact">
  <p class="branch-kicker">{esc(region)} BRANCHES</p><h1>{esc(region)} 와와학습코칭센터 지점안내</h1><p>{esc(description)}</p>
  <div class="branch-hero-actions"><a class="branch-button primary" href="#branch-list">{esc(region)} 센터 보기</a><a class="branch-button" href="/지점안내/">전국 지점안내</a></div>
</section>
<section class="branch-answer"><p class="branch-kicker">AREA SUMMARY</p><h2>{esc(region)} 지역 센터 구성</h2><p>{esc(district_copy)}으로 구성되어 있습니다. 주소와 수업 가능 동네를 함께 확인한 뒤 방문 전 상담에서 현재 시간표를 확인해 주세요.</p></section>
{coaching_system_section(f"{region} 지점", f"{region} 지역의 센터를 비교할 때 주소뿐 아니라 학생에게 필요한 플랜관리·학습관리·생활관리와 과목별 진단 흐름을 함께 확인해 보세요. 자료상 과목 운영 현황은 {subject_copy}입니다.")}
{search_panel(region)}
<section class="branch-section" id="branch-list"><div class="branch-section-head"><p class="branch-kicker">CENTERS</p><h2>{esc(region)} 센터 목록</h2><p data-branch-count>{len(centers)}개 센터를 표시하고 있습니다.</p></div><div class="branch-card-grid" data-branch-grid>{cards}</div><p class="branch-empty" data-branch-empty hidden>일치하는 센터가 없습니다.</p></section>
<section class="branch-section branch-faq" id="faq" aria-labelledby="region-faq-title"><div class="branch-section-head"><p class="branch-kicker">FAQ</p><h2 id="region-faq-title">{esc(region)} 지점 선택 질문</h2></div>{region_faq_html}</section>
<section class="branch-section branch-related"><div class="branch-section-head"><p class="branch-kicker">NEXT STEP</p><h2>센터 선택 후 확인할 내용</h2></div><div class="related-grid"><a href="/과목별학원/"><strong>과목별 학습 안내</strong><span>영어·수학·국어 등 과목별 관리 기준 보기</span></a><a href="/학년별학원/"><strong>학년별 학습 안내</strong><span>초등·중등·고등 단계별 학습 기준 보기</span></a><a href="/guide/parent-consultation-checklist/"><strong>상담 체크리스트</strong><span>상담 전 준비할 질문과 자료 확인하기</span></a></div></section>
'''
    html = finalize_page(upgrade_directory(page_shell(title, description, path, graph, "/assets/title.png", crumbs, body), centers, region), path)
    output = OUTPUT_ROOT / region / "index.html"
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(clean_html(html), encoding="utf-8")
    return path


def photo_gallery(center: dict[str, object]) -> str:
    photos = center["photos"]
    if not photos:
        return ""
    is_common = center.get("photoMode") == "common"
    intro = (
        "와와학습코칭센터의 학습 공간과 운영 환경을 살펴볼 수 있는 사진입니다. 센터별 공간 구성은 방문 전에 확인해 주세요."
        if is_common else
        "센터에서 제공한 학습 공간과 운영 사진입니다."
    )
    figures = []
    for index, photo in enumerate(photos, start=1):
        loading = "eager" if index == 1 else "lazy"
        priority = ' fetchpriority="high"' if index == 1 else ""
        alt = (
            f'{center["displayName"]} 학습 공간 참고 {index}'
            if is_common else
            f'{center["displayName"]} 학습 공간 {index}'
        )
        figures.append(
            f'<figure><img src="{esc(photo)}" alt="{esc(alt)}" width="1200" height="900" loading="{loading}" decoding="async"{priority}></figure>'
        )
    heading = f'{center["routeName"]} 학습 공간 살펴보기' if is_common else f'{center["routeName"]} 학습 공간'
    return f'''<section class="branch-media{' common-reference' if is_common else ''}" aria-labelledby="branch-media-title" data-photo-source="{'common' if is_common else 'center'}"><div class="branch-section-head"><p class="branch-kicker">LEARNING SPACE</p><h2 id="branch-media-title">{esc(heading)}</h2><p>{esc(intro)}</p></div><div class="branch-media-grid {'single' if len(photos) == 1 else ''}">{''.join(figures)}</div></section>'''


def primary_media(center: dict[str, object]) -> str:
    media = center["primaryMedia"]
    representative = media["representative"]
    body = media["body"]
    map_asset = media["map"]
    page_name = f'{center["displayName"]} 지점안내'
    avif_srcset = ", ".join(
        f'{item["src"]} {item["width"]}w' for item in body.get("variants", {}).get("avif", [])
    )
    webp_srcset = ", ".join(
        f'{item["src"]} {item["width"]}w' for item in body.get("variants", {}).get("webp", [])
    )
    sizes = "(max-width: 760px) calc(100vw - 32px), 760px"
    body_picture = (
        f'<picture>'
        f'<source type="image/avif" srcset="{esc(avif_srcset)}" sizes="{sizes}">'
        f'<source type="image/webp" srcset="{esc(webp_srcset)}" sizes="{sizes}">'
        f'<img src="{esc(body["src"])}" width="{body["width"]}" height="{body["height"]}" '
        f'alt="{esc(page_name)} 본문" loading="lazy" decoding="async" fetchpriority="low">'
        f'</picture>'
    )
    return f'''<section class="branch-primary-media" aria-label="{esc(page_name)} 본문 및 지도 이미지">
  <img class="branch-representative-image" src="{esc(representative["src"])}" width="{representative["width"]}" height="{representative["height"]}" alt="{esc(page_name)} 와와학습코칭센터 대표" style="display:none;" loading="lazy" decoding="async">
  <figure class="branch-body-image">{body_picture}<figcaption>{esc(center["region"])} 지역 학습코칭·수업 안내</figcaption></figure>
  <figure class="branch-map-image"><img src="{esc(map_asset["src"])}" width="{map_asset["width"]}" height="{map_asset["height"]}" alt="{esc(page_name)} 지도" loading="lazy" decoding="async"><figcaption>{esc(center["routeName"])} 위치 참고 지도 · 방문 전 주소와 운영 여부를 다시 확인해 주세요.</figcaption></figure>
</section>'''


def branch_topic_links(center: dict[str, object]) -> str:
    links = []
    for neighborhood in center["neighborhoods"]:
        for level, subject in BRANCH_TOPIC_VARIANTS:
            slug = f"{neighborhood}{level}{subject}학원"
            links.append(
                f'<a href="/지점안내/{esc(center["region"])}/{esc(center["routeName"])}/{esc(slug)}/">'
                f'<strong>{esc(neighborhood)} {esc(level)} {esc(subject)}학원</strong>'
                f'<span>{esc(level)} {esc(subject)} 학습 안내 보기</span></a>'
            )
    if not links:
        return ""
    return f'''<section class="branch-section branch-topic-links" id="learning-pages" aria-labelledby="learning-pages-title"><div class="branch-section-head"><p class="branch-kicker">LOCAL LEARNING PAGES</p><h2 id="learning-pages-title">동네별 영어·수학 학습 안내</h2><p>수업 가능 동네를 기준으로 학교급과 과목별 학습 준비 내용을 확인할 수 있습니다. 실제 개설 과목과 시간표는 상담에서 최종 확인해 주세요.</p></div><div class="branch-topic-link-grid">{''.join(links)}</div></section>'''


def school_section(center: dict[str, object]) -> str:
    cards = []
    for level in ("초등", "중등", "고등"):
        schools = center["schools"][level]
        pills = "".join(f"<li>{esc(school)}</li>" for school in schools) if schools else "<li>상담 시 확인</li>"
        cards.append(f'<article><h3>{level}학교</h3><ul>{pills}</ul></article>')
    neighborhoods = center["neighborhoods"]
    service = " · ".join(neighborhoods) if neighborhoods else f'{center["district"]} 인근'
    return f'''<section class="branch-section" id="schools" aria-labelledby="schools-title"><div class="branch-section-head"><p class="branch-kicker">SCHOOLS & AREA</p><h2 id="schools-title">인근 학교와 수업 가능 동네</h2><p>{esc(service)}을 기준으로 상담할 수 있습니다. 실제 통학 동선과 학교별 시험 범위는 상담에서 확인해 주세요.</p></div><div class="school-grid">{''.join(cards)}</div></section>'''


def subject_section(center: dict[str, object]) -> str:
    rows = []
    for subject in SUBJECTS:
        view = course_guidance(center, subject)
        notes = ''.join(f'<p class="branch-course-note">{esc(note)}</p>' for note in view["notes"])
        rows.append(f'<tr data-subject="{esc(subject)}"><th scope="row">{esc(subject)}</th><td><strong class="branch-grade-range">{esc(view["label"])}</strong>{notes}</td></tr>')
    conditions = ''.join(f'<p>{esc(note)}</p>' for note in center_notes(center))
    return f'''<section class="branch-section" id="subjects" aria-labelledby="subjects-title"><div class="branch-section-head"><p class="branch-kicker">SUBJECTS</p><h2 id="subjects-title">과목별 안내 학년과 수업 조건</h2><p>센터 운영 자료의 학년과 과목별 조건을 함께 정리했습니다. 문의로 표시된 범위는 수업 불가를 뜻하지 않으며, 개설 여부를 먼저 확인할 대상입니다.</p></div><div class="branch-table-wrap"><table><thead><tr><th>과목</th><th>안내 학년 · 확인할 조건</th></tr></thead><tbody>{''.join(rows)}</tbody></table></div><div class="branch-course-notes">{conditions}<p>{esc(REGISTRATION_NOTE)}</p></div></section>'''


def course_summary(center: dict[str, object]) -> str:
    cards = []
    for subject in ("영어", "수학"):
        view = course_guidance(center, subject)
        notes = ''.join(f'<p class="branch-course-note">{esc(note)}</p>' for note in view["notes"])
        cards.append(f'<div data-subject="{esc(subject)}"><dt>{esc(subject)}</dt><dd><strong class="branch-grade-range">{esc(view["label"])}</strong>{notes}</dd></div>')
    conditions = ''.join(f'<p>{esc(note)}</p>' for note in center_notes(center))
    return f'''<section class="branch-answer branch-course-summary" id="overview" aria-labelledby="overview-title"><p class="branch-kicker">QUICK ANSWER</p><h2 id="overview-title">{esc(center["routeName"])} 영어·수학 안내 학년</h2><dl class="branch-grade-grid">{''.join(cards)}</dl><div class="branch-course-notes">{conditions}<p>{esc(REGISTRATION_NOTE)}</p></div><a class="branch-inline-link" href="#subjects">국어·과학·사회와 전체 과목 조건 보기</a></section>'''


def learning_section(center: dict[str, object]) -> str:
    from branch_manuscript_editorial import learning_points

    points = learning_points(center)
    point_html = "".join(f'<article><h3>{esc(title)}</h3><p>{esc(copy)}</p></article>' for title, copy in points)
    areas = "·".join(center["neighborhoods"][:4]) or center["district"]
    return f'''<section class="branch-section" id="learning" aria-labelledby="learning-title"><div class="branch-section-head"><p class="branch-kicker">LEARNING GUIDE</p><h2 id="learning-title">{esc(center["routeName"])}에서 상담할 학습관리 항목</h2><p>{esc(areas)}에서 {esc(center["routeName"])}을 알아본다면, 학생의 최근 학교 자료와 집에서의 복습 기록을 함께 준비해 보세요. 아래 질문으로 필요한 도움을 구체화할 수 있습니다. 실제 제공 방식과 점검 주기는 상담에서 확인해 주세요.</p></div><div class="learning-grid">{point_html}</div><ol class="learning-flow"><li><b>1</b><span><strong>현재 상태 설명</strong>학교·학년과 최근 풀이에서 막힌 부분을 알려주세요.</span></li><li><b>2</b><span><strong>시작할 내용 질문</strong>현재 교재에서 먼저 보완할 단원과 이유를 물어보세요.</span></li><li><b>3</b><span><strong>복습 방법 확인</strong>집에서 혼자 해 볼 분량과 질문 전달 방법을 확인하세요.</span></li><li><b>4</b><span><strong>다음 점검 준비</strong>다시 풀어 본 기록 중 어떤 부분을 가져갈지 정해 보세요.</span></li></ol></section>'''


def generate_branch_page(center: dict[str, object]) -> str:
    from branch_page_summaries import center_summaries, validate_summaries

    region = center["region"]
    name = center["routeName"]
    path = f"/지점안내/{region}/{name}/"
    title = f'{center["displayName"]} | {region} {center["district"]} 지점안내'
    summary = center_summaries(center)
    validate_summaries(center, summary)
    description = summary["description"]
    crumbs = [("홈", "/"), ("지점안내", "/지점안내/"), (region, f"/지점안내/{region}/"), (name, path)]
    graph = graph_base(title, description, path, crumbs)
    page_url = encoded_url(path)
    media = center["primaryMedia"]
    image = media["representative"]["src"]
    image_urls = [
        DOMAIN + quote(media[kind]["src"], safe="/")
        for kind in ("representative", "body", "map")
    ]
    views = [course_guidance(center, subject) for subject in SUBJECTS]
    weekend = weekend_guidance(center)
    offers = [
        {"@type": "Offer", "url": page_url, "itemOffered": {
            "@type": "Service", "name": f'{name} {view["subject"]} 학습코칭', "serviceType": "학습코칭",
            "description": " ".join([f'{view["subject"]} 안내 학년: {view["label"]}.', *view["notes"], *center_notes(center), REGISTRATION_NOTE]),
            "audience": {"@type": "EducationalAudience", "educationalRole": "student", "audienceType": view["label"]},
        }}
        for view in views if view["grades"]
    ] or [
        {"@type": "Offer", "itemOffered": {"@type": "Service", "name": f'{center["routeName"]} 학습 상담', "serviceType": "학습코칭 상담"}}
    ]

    postal = {"@type": "PostalAddress", "streetAddress": center["address"], "addressRegion": region, "addressLocality": center["district"], "addressCountry": "KR"}
    organization = {
        "@type": ["EducationalOrganization", "LocalBusiness"],
        "@id": page_url + "#academy",
        "name": center["displayName"],
        "legalName": center["registeredName"],
        "url": page_url,
        "image": image_urls,
        "address": postal,
        "parentOrganization": {"@id": DOMAIN + "/#organization"},
        "areaServed": [{"@type": "Place", "name": value} for value in center["neighborhoods"][:12]],
        "knowsAbout": [f'{view["subject"]} 학습코칭' for view in views if view["grades"]] or ["학습 상담"],
        "identifier": center["registrationNumber"],
        "additionalProperty": {
            "@type": "PropertyValue",
            "name": "센터 정보 확인 기준일",
            "value": center["informationReviewedAt"],
        },
        "makesOffer": offers,
    }
    graph.append(organization)
    webpage = next(item for item in graph if item.get("@id") == page_url + "#webpage")
    webpage["mainEntity"] = {"@id": page_url + "#academy"}
    webpage["about"] = [{"@id": page_url + "#academy"}, {"@type": "Place", "name": f'{region} {center["district"]}'}]
    webpage["primaryImageOfPage"] = {"@type": "ImageObject", "url": image_urls[0]}
    webpage["hasPart"] = [
        {"@type": "WebPageElement", "@id": page_url + "#center-info", "name": f"{name} 기본정보"},
        {"@type": "WebPageElement", "@id": page_url + "#overview", "name": f"{name} 영어·수학 안내 학년"},
        {"@type": "WebPageElement", "@id": page_url + "#subjects", "name": "과목별 안내 학년과 수업 조건"},
        {"@type": "WebPageElement", "@id": page_url + "#schools", "name": "인근 학교와 수업 가능 동네"},
        {"@type": "WebPageElement", "@id": page_url + "#learning", "name": "학습관리 항목"},
        {"@type": "WebPageElement", "@id": page_url + "#faq", "name": f"{name} 자주 묻는 질문"},
    ]
    child_items = []
    for neighborhood in center["neighborhoods"]:
        for level, subject in BRANCH_TOPIC_VARIANTS:
            child_path = f'/지점안내/{region}/{name}/{neighborhood}{level}{subject}학원/'
            child_items.append({
                "@type": "ListItem",
                "position": len(child_items) + 1,
                "name": f"{neighborhood} {level} {subject}학원",
                "url": encoded_url(child_path),
            })
    if child_items:
        webpage["hasPart"].append({"@id": page_url + "#learning-pages"})
        graph.append({
            "@type": "ItemList",
            "@id": page_url + "#learning-pages",
            "name": f"{name} 동네별 영어·수학 학습 안내",
            "numberOfItems": len(child_items),
            "itemListElement": child_items,
        })
    graph.append({
        "@type": "Article", "@id": page_url + "#article", "headline": title,
        "description": description, "abstract": summary["abstract"], "inLanguage": "ko-KR",
        **page_dates(path), "mainEntityOfPage": {"@id": page_url + "#webpage"},
        "about": {"@id": page_url + "#academy"}, "articleSection": ["대표·본문·지도 이미지", "학습 공간", "센터 기본정보", "가능 과목과 학년", "인근 학교", "학습관리", "상담 안내"],
        "image": image_urls,
        "author": {"@id": DOMAIN + "/#organization"},
        "publisher": {"@id": DOMAIN + "/#organization"},
    })
    graph.append({
        "@type": "Service", "@id": page_url + "#service",
        "name": f'{center["displayName"]} 학습코칭 상담',
        "serviceType": "초중고 학습코칭",
        "description": f'{center["routeName"]}에서 학생의 학교·학년·현재 교재와 과목별 어려움을 확인하고 학습 계획을 상담합니다.',
        "provider": {"@id": page_url + "#academy"},
        "areaServed": [{"@type": "Place", "name": value} for value in center["neighborhoods"][:12]],
        "audience": {"@type": "EducationalAudience", "educationalRole": "student"},
    })

    neighborhoods = " · ".join(center["neighborhoods"][:8]) if center["neighborhoods"] else f'{center["district"]} 인근'
    faq_items = [
        (f'{name}은 어디에 있나요?', f'{center["displayName"]}의 주소는 {center["address"]}입니다. 방문 전 상담으로 건물과 입실 방법을 다시 확인해 주세요.'),
        (f'{name}에서 어떤 과목과 학년을 상담할 수 있나요?', branch_course_answer(center)),
        (f'{name}의 수업 가능 동네와 인근 학교는 어디인가요?', f'{neighborhoods}을 중심으로 상담할 수 있습니다. 학교별 시험 범위와 통학 가능 여부는 학생의 학교를 알려주고 확인해 주세요.'),
        (f'{name}은 주말에도 수업하나요?', weekend),
        ("상담 전에 무엇을 준비하면 좋나요?", "학생의 학년과 학교, 최근 시험지, 현재 교재, 어려운 단원, 숙제 수행 정도와 평소 공부 시간을 정리해 오면 학습 계획을 구체적으로 확인하기 좋습니다."),
    ]
    graph.append({"@type": "FAQPage", "@id": page_url + "#faq", "mainEntity": [
        {"@type": "Question", "name": q, "acceptedAnswer": {"@type": "Answer", "text": a}} for q, a in faq_items
    ]})

    map_query = quote(center["address"])
    location = center["locationGuide"] or "방문 전 정확한 건물과 층수, 입실 방법을 상담에서 다시 확인해 주세요."
    neighborhood_pills = "".join(f"<li>{esc(value)}</li>" for value in center["neighborhoods"]) or f"<li>{esc(center['district'])} 인근</li>"
    faq_html = "".join(f'<details><summary>{esc(q)}</summary><p>{esc(a)}</p></details>' for q, a in faq_items)
    related = [
        (f"{region} 지점안내", f"/지점안내/{region}/", "같은 지역의 다른 센터 보기"),
        ("과목별학원", "/과목별학원/", "과목별 학습관리 기준 확인"),
        ("학년별학원", "/학년별학원/", "학년별 학습 준비 확인"),
        ("학습가이드", "/guide/", "플래너·오답·시험 준비 가이드"),
    ]
    related_html = "".join(f'<a href="{href}"><strong>{esc(label)}</strong><span>{esc(copy)}</span></a>' for label, href, copy in related)

    body = f'''
<section class="branch-detail-hero">
  <div><p class="branch-kicker">{esc(region)} · {esc(center["district"])} BRANCH</p><h1>{esc(center["displayName"])}</h1><p class="branch-lead">{esc(summary["lead"])}</p></div>
  <dl class="hero-facts"><div><dt>주소</dt><dd>{esc(center["address"])}</dd></div><div><dt>과목·학년</dt><dd><a href="#subjects">과목별 학년과 수업 조건 확인</a></dd></div><div><dt>수업 가능 동네</dt><dd>{esc(neighborhoods)}</dd></div></dl>
  <div class="branch-hero-actions"><a class="branch-button primary" href="#overview">학년 먼저 확인</a><a class="branch-button" href="#center-info">센터 정보</a><a class="branch-button" href="#consult">상담 준비</a></div>
</section>
{course_summary(center)}
{primary_media(center)}
{photo_gallery(center)}
<nav class="branch-toc" aria-label="페이지 목차"><strong>목차</strong><a href="#center-info">기본정보</a><a href="#subjects">과목·학년</a><a href="#schools">학교·동네</a><a href="#learning">학습관리</a><a href="#faq">자주 묻는 질문</a></nav>
<section class="branch-section" id="center-info" aria-labelledby="center-info-title"><div class="branch-section-head"><p class="branch-kicker">CENTER INFORMATION</p><h2 id="center-info-title">{esc(name)} 기본정보</h2><p>센터에서 제공한 등록 정보와 운영 자료를 기준으로 정리했습니다.</p></div><dl class="info-list"><div><dt>주소</dt><dd>{esc(center["address"])}</dd></div><div><dt>등록 명칭</dt><dd>{esc(center["registeredName"])}</dd></div><div><dt>교육지원청 등록번호</dt><dd>{esc(center["registrationNumber"])}</dd></div>{f'<div><dt>등록일</dt><dd>{esc(center["registrationDate"])}</dd></div>' if center["registrationDate"] else ''}<div><dt>정보 확인 기준일</dt><dd>{esc(center["informationReviewedAt"])} · 제공된 센터 자료 기준</dd></div><div><dt>평일 운영 참고</dt><dd>{esc(center["openingReference"])} · 실제 방문·수업 시간은 상담 확인</dd></div><div><dt>주말 운영</dt><dd>{esc(weekend)}</dd></div><div><dt>위치 안내</dt><dd>{esc(location)}</dd></div></dl><div class="info-actions"><a class="branch-button primary" href="https://map.naver.com/p/search/{map_query}" target="_blank" rel="noopener noreferrer">네이버 지도에서 주소 검색</a><a class="branch-button" href="tel:{PHONE_LINK}">전화 상담</a></div></section>
{subject_section(center)}
{school_section(center)}
<section class="branch-section service-area" aria-labelledby="service-area-title"><div class="branch-section-head"><p class="branch-kicker">SERVICE AREA</p><h2 id="service-area-title">수업 가능 동네</h2><p>센터 상담 자료에 연결된 동네를 정리했습니다. 실제 통학 거리와 시간표는 주소를 기준으로 다시 확인해 주세요.</p></div><ul>{neighborhood_pills}</ul></section>
{learning_section(center)}
<section class="branch-section consult-section" id="consult" aria-labelledby="consult-title"><div class="branch-section-head"><p class="branch-kicker">CONSULTATION</p><h2 id="consult-title">상담 전에 준비하면 좋은 내용</h2><p>학생의 현재 상태를 구체적으로 알려주면 시작 단원과 관리 우선순위를 더 정확하게 확인할 수 있습니다.</p></div><ul class="consult-checklist"><li>학교와 학년, 최근 시험 점수와 시험지</li><li>현재 사용하는 교재와 공부 중인 단원</li><li>자주 틀리는 문제 유형과 어려운 과목</li><li>숙제 수행 정도와 평소 공부 시간</li><li>희망 과목, 가능한 요일과 시간</li><li>선행과 결손 보충 중 우선할 목표</li></ul><div class="consult-cta"><div><strong>{esc(center["displayName"])} 상담</strong><p>현재 운영 과목과 시간표는 상담에서 최종 확인해 주세요.</p></div><a class="branch-button primary" href="{CONSULT_URL}" target="_blank" rel="noopener noreferrer">상담 신청</a></div></section>
<section class="branch-section branch-faq" id="faq" aria-labelledby="faq-title"><div class="branch-section-head"><p class="branch-kicker">FAQ</p><h2 id="faq-title">{esc(name)} 자주 묻는 질문</h2></div>{faq_html}</section>
{branch_topic_links(center)}
<section class="branch-section branch-related" aria-labelledby="related-title"><div class="branch-section-head"><p class="branch-kicker">RELATED GUIDE</p><h2 id="related-title">함께 확인할 안내</h2></div><div class="related-grid">{related_html}</div></section>
'''
    html = finalize_page(upgrade_center(page_shell(title, description, path, graph, image, crumbs, body), center), path)
    output = OUTPUT_ROOT / region / name / "index.html"
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(clean_html(html), encoding="utf-8")
    return path


def update_sitemap(paths: list[str]) -> None:
    sitemap = ROOT / "sitemap.xml"
    text = sitemap.read_text(encoding="utf-8")
    start = "  <!-- branch-directory:start -->"
    end = "  <!-- branch-directory:end -->"
    pattern = r"^[ \t]*" + re.escape(start.strip()) + r".*?" + re.escape(end.strip()) + r"[ \t]*(?:\r?\n|$)"
    text = re.sub(pattern, "", text, flags=re.S | re.M)
    entries = []
    for path in paths:
        entries.append(f"  <url>\n    <loc>{encoded_url(path)}</loc>\n    <lastmod>{page_dates(path)['dateModified']}</lastmod>\n  </url>")
    block = start + "\n" + "\n".join(entries) + "\n" + end + "\n"
    text = text.replace("</urlset>", block + "</urlset>")
    sitemap.write_text(text, encoding="utf-8")


def write_manifest(centers: list[dict[str, object]], excluded: list[dict[str, str]]) -> None:
    DATA_ROOT.mkdir(parents=True, exist_ok=True)
    public = []
    for center in centers:
        item = {key: value for key, value in center.items() if key not in {"photoSource"}}
        public.append(item)
    payload = {
        "generatedAt": TODAY,
        "sources": [SOURCE_WORKBOOK.name, TARGET_WORKBOOK.name, PHOTO_ROOT.name, SUPPLEMENTAL_CENTERS_FILE.name],
        "routePattern": "/지점안내/{지역}/{지점명}/",
        "centers": public,
        "excluded": excluded,
    }
    (DATA_ROOT / "branches.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def audit(centers: list[dict[str, object]], paths: list[str], excluded: list[dict[str, str]]) -> dict[str, object]:
    canonical_re = re.compile(r'<link rel="canonical" href="([^"]+)">')
    h1_re = re.compile(r"<h1[^>]*>(.*?)</h1>", re.S)
    missing = []
    canonicals = []
    titles = []
    for path in paths:
        relative = path.strip("/").split("/") if path.strip("/") else []
        file = ROOT.joinpath(*relative, "index.html")
        if not file.exists():
            missing.append(path)
            continue
        text = file.read_text(encoding="utf-8")
        canonical = canonical_re.search(text)
        h1 = h1_re.search(text)
        canonicals.append(canonical.group(1) if canonical else "")
        titles.append(re.sub("<[^>]+>", "", h1.group(1)).strip() if h1 else "")
    broken_images = []
    for center in centers:
        primary_urls = [center["primaryMedia"][kind]["src"] for kind in ("representative", "body", "map")]
        for url in [*center["photos"], *primary_urls]:
            if not (ROOT / url.lstrip("/")).exists():
                broken_images.append(url)
    report = {
        "generatedAt": TODAY,
        "centerCount": len(centers),
        "regionCount": len({center["region"] for center in centers}),
        "regionCounts": dict(Counter(center["region"] for center in centers)),
        "pageCount": len(paths),
        "withPhotos": sum(bool(center["photos"]) for center in centers),
        "withoutPhotos": [center["routeName"] for center in centers if not center["photos"]],
        "withCenterPhotos": sum(center.get("photoMode") == "center" for center in centers),
        "withCommonPhotoFallback": sum(center.get("photoMode") == "common" for center in centers),
        "withPrimaryMedia": sum(bool(center.get("primaryMedia")) for center in centers),
        "verifiedMapMatches": sum(center.get("primaryMedia", {}).get("mapMatch") in {"verified-center-source-row", "verified-reference-center"} for center in centers),
        "supplementalCenters": [center["routeName"] for center in centers if center.get("verifiedMediaKey")],
        "missingTargetMapping": [center["routeName"] for center in centers if not center["neighborhoods"]],
        "excluded": excluded,
        "missingFiles": missing,
        "duplicateCanonicals": [value for value, count in Counter(canonicals).items() if value and count > 1],
        "missingCanonicalCount": sum(not value for value in canonicals),
        "missingH1Count": sum(not value for value in titles),
        "brokenImages": broken_images,
    }
    REPORT_ROOT.mkdir(parents=True, exist_ok=True)
    (REPORT_ROOT / "audit.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    if missing or report["duplicateCanonicals"] or report["missingCanonicalCount"] or report["missingH1Count"] or broken_images:
        raise RuntimeError("지점안내 감사 실패: reports/branch-directory/audit.json 확인")
    return report


def main() -> None:
    for required in (
        SOURCE_WORKBOOK, TARGET_WORKBOOK, PHOTO_ROOT, COMMON_PHOTO_ROOT,
        REPRESENTATIVE_ROOT, BODY_IMAGE_ROOT, MAP_IMAGE_ROOT, VERIFIED_MEDIA_MANIFEST, SUPPLEMENTAL_CENTERS_FILE,
    ):
        if not required.exists():
            raise FileNotFoundError(required)
    centers, excluded = load_centers()
    if len(centers) != 193:
        raise RuntimeError(f"검토된 센터 수가 예상과 다릅니다: {len(centers)} (예상 193)")

    import_media(centers)
    import_primary_media(centers)
    if OUTPUT_ROOT.exists():
        shutil.rmtree(OUTPUT_ROOT)

    paths = [generate_hub(centers)]
    for region in REGIONS:
        region_centers = [center for center in centers if center["region"] == region]
        if region_centers:
            paths.append(generate_region_page(region, region_centers))
            paths.extend(generate_branch_page(center) for center in region_centers)

    update_sitemap(paths)
    save_dates()
    write_manifest(centers, excluded)
    report = audit(centers, paths, excluded)
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
